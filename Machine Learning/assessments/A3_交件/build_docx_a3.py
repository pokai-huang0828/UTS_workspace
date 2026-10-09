"""把 A3_報告.md 的「報告文字」轉成 Word（沿用 A2 build_docx.py：去掉私註、要點、圖的規格行）。

用法：python build_docx_a3.py → Huang_26254793_321513_A3.docx
接著：powershell -ExecutionPolicy Bypass -File to_pdf_a3.ps1（Word 更新目錄、匯出 PDF）

與 A2 版的差異：SRC／OUT／圖檔對照、各表欄寬、圖號可為 B1–B3（附錄）、儲存格內 <br> 換段、
表 2「判定」欄依文字加淡色底（文字為準，顏色只輔助）、建置後掃描 docx 有無私註殘留、
嵌入圖時裁掉圖頂端的英文標題列（Notebook 的 Fig N 編號；報告圖號與對照寫在中文圖說）。
"""
import re
import sys
from io import BytesIO
from pathlib import Path

import numpy as np
from PIL import Image

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "A3_報告.md"
OUT = HERE / "Huang_26254793_321513_A3.docx"
FIGDIR = HERE / "figs"
# 報告圖號 → Notebook 輸出的圖檔（圖說中另註 Notebook Fig 編號）
# HD 修訂（2026-10-09）：圖 1＝Nb Fig 8（每輪的錢）、圖 2＝Nb Fig 9（AUC 拆解）、圖 3＝Nb Fig 1、圖 4＝Nb Fig 7、
# 圖 5＝Nb Fig 3、圖 6＝Nb Fig 4。圖寬以「圖內 9 pt 字嵌入後約 ≥ 7.5 pt」與版面為準（原圖 150 dpi，寬 14.6–17.8 cm）。
FIGS = {"1": FIGDIR / "fig8_round_spend.png", "2": FIGDIR / "fig9_auc_decomposition.png",
        "3": FIGDIR / "fig1_gains.png", "4": FIGDIR / "fig7_round_cost.png",
        "5": FIGDIR / "fig3_uplift.png", "6": FIGDIR / "fig4_segment_auc.png",
        "B1": FIGDIR / "fig2_cost_vs_depth.png", "B2": FIGDIR / "fig6_confusion_matrix.png",
        "B3": FIGDIR / "fig5_perm_importance.png"}
# hd-fix（2026-10-09）：依 render 審查放寬圖寬（版心 15.9 cm），讓圖內 9 pt 字嵌入後約 ≥ 8 pt；圖 3 較高，略窄以免整頁留白
FIG_W = {"1": 15.9, "2": 15.9, "3": 14.6, "4": 15.9, "5": 14.6, "6": 14.8,
         "B1": 14.9, "B2": 15.6, "B3": 15.9}  # cm
FIG_AFTER_HEADING = set()

LATIN = "Times New Roman"
CJK = "新細明體"
BODY_PT = 12
SMALL_PT = 10


# ---------------------------------------------------------------- 字型與樣式
def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = LATIN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), LATIN)
    rfonts.set(qn("w:eastAsia"), CJK)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_fonts(style, size, bold=False, color="000000"):
    style.font.name = LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), LATIN)
    rfonts.set(qn("w:eastAsia"), CJK)


def para_format(p, spacing=2.0, before=0, after=0, align=None, first_indent=None):
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if first_indent is not None:
        pf.first_line_indent = first_indent


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def add_inline(p, text, size=BODY_PT, color=None, bold=None):
    """**粗體**、*斜體* 轉成 run。"""
    text = text.replace("<br>", " ")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_run_font(p.add_run(part[2:-2]), size, bold=True, color=color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_run_font(p.add_run(part[1:-1]), size, italic=True, color=color)
        else:
            set_run_font(p.add_run(part), size, bold=bold, color=color)


def cell_text(cell, txt, size, spacing=1.1, after=1, bold=None, keep=False):
    """儲存格文字；<br> 換成新段落。"""
    cell.text = ""
    for pi, part in enumerate([s.strip() for s in txt.split("<br>")]):
        p = cell.paragraphs[0] if pi == 0 else cell.add_paragraph()
        para_format(p, spacing, after=after)
        if keep:
            p.paragraph_format.keep_with_next = True
        add_inline(p, part, size=size, bold=bold)


def shade(cell, hexcolor):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_field(p, instr, placeholder="1"):
    r = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = placeholder
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, t, s):
        r._element.append(el)
    r2 = p.add_run(); r2._element.append(txt)
    r3 = p.add_run(); r3._element.append(e)
    set_run_font(r2, SMALL_PT, color="666666")


# ---------------------------------------------------------------- 解析 Markdown → 報告區塊
def parse(md):
    lines = md.split("\n")
    blocks = []  # (kind, payload)
    i = 0
    skip_private = False
    in_figspec = None
    while i < len(lines) and not lines[i].startswith("## 封面"):
        i += 1
    while i < len(lines):
        ln = lines[i]
        if re.match(r"#{3,4} (要點|交給第 5 節)", ln):
            skip_private = True
            i += 1
            continue
        if re.match(r"#{2,4} ", ln):
            skip_private = False
            in_figspec = None
            if re.match(r"#{3,4} (正文|封面內容)$", ln):
                i += 1
                continue
            level = 1 if ln.startswith("## ") else 2
            title = re.sub(r"^#+ ", "", ln)
            title = re.sub(r"（(Appendix|Cover|References)[^）]*）$", "", title)
            blocks.append(("h", (level, title)))
            i += 1
            continue
        if skip_private or ln.startswith(">") or ln.strip() == "---":
            i += 1
            continue
        m = re.match(r"\*\*圖 ([0-9B]+)｜", ln)
        if m:
            in_figspec = m.group(1)
            i += 1
            continue
        if in_figspec:
            cap = re.match(r"- \*\*圖說\*\*：(.*)", ln)
            if cap:
                blocks.append(("fig", (in_figspec, cap.group(1).strip())))
                in_figspec = None
            i += 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
            blocks.append(("table", cells))
            continue
        if ln.startswith("<p align=\"center\">"):
            blocks.append(("eq", re.sub(r"</?p[^>]*>", "", ln)))
            i += 1
            continue
        m = re.match(r"\*\*框｜(.+)\*\*$", ln)
        if m:
            blocks.append(("boxcap", m.group(1)))
            i += 1
            continue
        if re.match(r"\*\*表 R?[0-9A]+｜", ln):
            blocks.append(("tcap", ln))
            i += 1
            continue
        if (ln.startswith("註：") or ln.startswith("註 ") or re.match(r"^\**表 R?[0-9A]+ 註", ln)
                or (re.match(r"^[①②③④⑤]", ln) and blocks and blocks[-1][0] == "note")):
            blocks.append(("note", ln))
            i += 1
            continue
        if ln.startswith("- "):
            blocks.append(("bullet", ln[2:]))
            i += 1
            continue
        if ln.strip():
            blocks.append(("p", ln.strip()))
        i += 1
    return blocks


# ---------------------------------------------------------------- 產生文件
# 各表欄寬（cm，合計約 15.9）
# hd-fix：表號改為 R1–R7、RA（與 A2 表號區隔）；表 R7 刪「A2 節」欄
WIDTHS = {"R1": [2.9, 5.1, 4.4, 3.5],            # Notebook 補全、選模、發現
          "R2": [2.3, 6.1, 1.6, 5.9],            # 三個商業目標
          "R3": [2.8, 2.8, 5.1, 2.5, 2.7],       # 十條成功標準
          "R4": [2.1, 8.0, 2.0, 3.8],            # A2 表 8 其餘面向
          "R5": [2.4, 3.0, 2.0, 2.6, 2.6, 3.3],  # 三種錯誤
          "R6": [2.6, 1.7, 3.4, 2.9, 2.9, 2.4],  # 改進、新增功能、新產品
          "R7": [1.5, 5.4, 3.1, 4.1, 1.8],       # 後續問題步驟
          "RA": [2.8, 3.2, 2.8, 3.3, 3.8]}
VERDICT_SHADE = [("達標", "E2EEDD"), ("未達", "F5DCD8"), ("未證實", "FAEED2"), ("部分", "FAEED2"),
                 ("示算", "ECEBE8"), ("列報", "ECEBE8"), ("不可離線驗", "ECEBE8"), ("量級足夠", "ECEBE8")]
NOTE_PT = 10
CAP_PT = 11
BOX_W = [3.4, 12.5]


def fig_stream(path):
    """裁掉圖頂端的標題列：找第一段深色列（標題），從其後第一條全白列開始保留；找不到就原圖。"""
    im = Image.open(path).convert("RGB")
    dark = (np.asarray(im.convert("L")) < 200).any(axis=1)
    rows = np.flatnonzero(dark)
    if len(rows):
        r0 = int(rows[0])
        r1 = r0
        while r1 < len(dark) and dark[r1]:
            r1 += 1
        if r0 < 60 and r1 - r0 < 40 and r1 < len(dark) and not dark[r1]:
            im = im.crop((0, r1, im.width, im.height))
    buf = BytesIO()
    im.save(buf, format="PNG", dpi=(150, 150))
    buf.seek(0)
    return buf


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def set_col_widths(t, widths):
    t.autofit = False
    for ci, w in enumerate(widths):
        if ci < len(t.columns):
            t.columns[ci].width = Cm(w)
    for row in t.rows:
        for ci, w in enumerate(widths):
            if ci < len(row.cells):
                row.cells[ci].width = Cm(w)


def add_box(doc, title, rows, spacer=True, keep=True):
    """「框」＝兩欄表格：合併的標題列（深底）＋每列「項目｜內容」（略過 md 表頭列）。
    keep=True：整框不跨頁（決策摘要）；False：列不拆開、但框可跨頁（標題列重複），避免整框推到下一頁留白。"""
    t = doc.add_table(rows=0, cols=2)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    head = t.add_row().cells
    cell = head[0].merge(head[1])
    p = cell.paragraphs[0]
    para_format(p, 1.1, before=2, after=2)
    add_inline(p, "**" + title + "**", size=NOTE_PT + 0.5)
    shade(cell, "D9D6CF")
    for row in rows:
        cells = t.add_row().cells
        for ci, txt in enumerate((row + ["", ""])[:2]):
            cell_text(cells[ci], txt, NOTE_PT, spacing=1.15, bold=True if ci == 0 else None)
        shade(cells[0], "F3F2EE")
    for ci, w in enumerate(BOX_W):
        t.columns[ci].width = Cm(w)
    cell.width = Cm(sum(BOX_W))
    for r in t.rows[1:]:
        for ci, w in enumerate(BOX_W):
            r.cells[ci].width = Cm(w)
    for ri, r in enumerate(t.rows):  # 列不拆開；keep 時整框不跨頁，否則只把標題列綁住第一列
        no_split(r)
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = (ri < len(t.rows) - 1) if keep else ri == 0
    repeat_header(t.rows[0])
    if not spacer:  # 下一個是標題：不加間隔段，避免框剛好填滿一頁時多出一頁空白
        return
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    sp.paragraph_format.line_spacing = Pt(2)


def build():
    from docx.enum.text import WD_TAB_ALIGNMENT

    md = SRC.read_text(encoding="utf-8")
    blocks = parse(md)
    doc = Document()
    cp = doc.core_properties
    cp.author = "黃柏凱 (Po-Kai Huang)"
    cp.last_modified_by = "Po-Kai Huang"
    cp.comments = ""
    cp.title = "Bank X 信用卡外撥名單：從概念驗證到商業評估"
    cp.subject = "321513 Machine Learning - Assessment Task 3 (Business focus)"

    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))

    style_fonts(doc.styles["Normal"], BODY_PT)
    op = OxmlElement("w:overflowPunct")
    op.set(qn("w:val"), "0")
    doc.styles["Normal"].element.get_or_add_pPr().append(op)
    style_fonts(doc.styles["Heading 1"], 14, bold=True)
    style_fonts(doc.styles["Heading 2"], 13, bold=True)
    for name in ("Heading 1", "Heading 2"):
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(12), Pt(6), 1.5
        pf.keep_with_next = True

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE", "1")

    # ---- 封面
    assert blocks[0] == ("h", (1, "封面")), blocks[0]
    idx = 1
    cover_lines, cover_table = [], None
    while blocks[idx][0] != "h":
        k, v = blocks[idx]
        if k == "p":
            cover_lines.append(v)
        elif k == "table":
            cover_table = v
        idx += 1
    for _ in range(5):
        para_format(doc.add_paragraph(), 1.0)
    for j, t in enumerate(cover_lines):
        p = doc.add_paragraph()
        para_format(p, 1.3, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_inline(p, t, size=(20 if j == 0 else 15 if j == 1 else 12))
    para_format(doc.add_paragraph(), 1.0, after=24)
    tbl = doc.add_table(rows=0, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in cover_table[1:]:
        r = tbl.add_row().cells
        for ci, (c, txt) in enumerate(zip(r, row)):
            cell_text(c, txt, 11, spacing=1.3, after=4, bold=True if ci == 0 else None)
    set_col_widths(tbl, [3.2, 12.0])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- 目錄
    p = doc.add_paragraph()
    para_format(p, 1.5, after=12)
    add_inline(p, "**目錄**", size=14)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "（目錄）")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- 內文
    in_refs = False
    cur_table = None
    pending_note = None
    pending_box = None
    body = blocks[idx:]
    for bi, (k, v) in enumerate(body):
        nxt = body[bi + 1][0] if bi + 1 < len(body) else None
        if k == "h":
            level, title = v
            in_refs = title.startswith("參考文獻")
            h = doc.add_heading(level=level)
            if title.startswith("參考文獻") or title.startswith("附錄 A"):
                h.paragraph_format.page_break_before = True  # 不用分頁段落，避免前一頁剛好滿時多出空白頁
            add_inline(h, title, size=14 if level == 1 else 13)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
        elif k == "p":
            p = doc.add_paragraph()
            if in_refs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.line_spacing = Pt(30)
                p.paragraph_format.left_indent = Cm(1.27)
                p.paragraph_format.first_line_indent = Cm(-1.27)
                add_inline(p, v, size=BODY_PT)
            else:
                para_format(p, 2.0, after=0)
                add_inline(p, v)
        elif k == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            para_format(p, 2.0, after=0)
            add_inline(p, v, size=BODY_PT)
        elif k == "eq":
            p = doc.add_paragraph()
            para_format(p, 1.5, before=4, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_inline(p, v)
        elif k == "tcap":
            m = re.match(r"(\*\*表 (R?[0-9A]+)｜[^*]+\*\*)(.*)$", v)
            cur_table = m.group(2) if m else None
            title, rest = (m.group(1), m.group(3).strip()) if m else (v, "")
            if rest.startswith("（") and rest.endswith("）"):
                pending_note = "註：" + rest[1:-1]
                rest = ""
            p = doc.add_paragraph()
            para_format(p, 1.2, before=10, after=4)
            p.paragraph_format.keep_with_next = True
            add_inline(p, title + rest, size=CAP_PT)
        elif k == "note":
            p = doc.add_paragraph()
            para_format(p, 1.2, before=3, after=8)
            p.paragraph_format.keep_together = True
            add_inline(p, v, size=NOTE_PT, color="333333")
        elif k == "boxcap":
            pending_box = v
        elif k == "table" and pending_box:
            add_box(doc, pending_box, v[1:], spacer=(nxt != "h"), keep=pending_box.startswith("決策摘要"))
            pending_box = None
        elif k == "table":
            ncol = len(v[0])
            size = 9 if ncol < 6 else 8.5
            t = doc.add_table(rows=0, cols=ncol)
            t.style = doc.styles["Table Grid"]
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            verdict_col = v[0].index("判定") if "判定" in v[0] else None
            for ri, row in enumerate(v):
                cells = t.add_row().cells
                no_split(t.rows[-1])
                row = (row + [""] * ncol)[:ncol]
                for ci, (c, txt) in enumerate(zip(cells, row)):
                    cell_text(c, txt, size, bold=True if ri == 0 else None, keep=(ri == 0))
                    if ri == 0:
                        shade(c, "E8E6E1")
                    elif ci == verdict_col:
                        plain = txt.replace("**", "").strip()
                        for key, col in VERDICT_SHADE:
                            if plain.startswith(key):
                                shade(c, col)
                                break
                if ri == 0:
                    repeat_header(t.rows[0])
            if cur_table in WIDTHS and len(WIDTHS[cur_table]) == ncol:
                set_col_widths(t, WIDTHS[cur_table])
            elif cur_table is not None:
                print(f"WARNING: 表 {cur_table} 欄寬未設定或欄數不符（{ncol} 欄）", file=sys.stderr)
            cur_table = None
            if pending_note:
                p = doc.add_paragraph()
                para_format(p, 1.2, before=3, after=8)
                add_inline(p, pending_note, size=NOTE_PT, color="333333")
                pending_note = None
            else:
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif k == "fig":
            num, cap = v
            p = doc.add_paragraph()
            para_format(p, 1.0, before=8, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(fig_stream(FIGS[num]), width=Cm(FIG_W[num]))
            c = doc.add_paragraph()
            para_format(c, 1.2, after=10)
            c.paragraph_format.keep_together = True  # hd-fix：圖說不跨頁（圖段已 keep_with_next）
            add_inline(c, cap.replace(f"圖 {num}｜", f"**圖 {num}｜**", 1), size=CAP_PT)
    doc.save(OUT)
    return OUT, blocks


def scan(path):
    """建置後掃描：私註、要點、警示符號、未換掉的 md 標記不得殘留。"""
    d = Document(str(path))
    texts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                texts.append(c.text)
    bad = ("私註", "要點", "⚠", "交給第 5 節", "**", "<br>", "TODO:", "[待驗]")
    hits = [(b, s[:60]) for s in texts for b in bad if b in s]
    return hits


if __name__ == "__main__":
    out, blocks = build()
    kinds = {}
    for k, _ in blocks:
        kinds[k] = kinds.get(k, 0) + 1
    print(out, kinds)
    hits = scan(out)
    if hits:
        print("RESIDUE FOUND:", hits)
        sys.exit(1)
    print("scan: clean")
