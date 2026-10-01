"""把 A2_草稿_v1.md 的「報告文字」轉成 Word（去掉要點表、私註、字數行、圖的規格清單）。

用法：python build_docx.py → Huang_26254793_321513_A2.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE.parent / "A2_草稿_v1.md"
OUT = HERE / "Huang_26254793_321513_A2.docx"
FIGS = {"1": HERE / "fig1_cost_vs_leakage.png", "2": HERE / "fig2_mlops_level1.png"}

LATIN = "Times New Roman"
CJK = "新細明體"
BODY_PT = 12
SMALL_PT = 10


# ---------------------------------------------------------------- 字型與樣式
def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = LATIN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), LATIN)
    rfonts.set(qn("w:eastAsia"), CJK)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_fonts(style, size, bold=False, color="000000"):
    style.font.name = LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), LATIN)
    rfonts.set(qn("w:eastAsia"), CJK)


def para_format(p, spacing=2.0, before=0, after=0, align=None, first_indent=None):
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if first_indent is not None:
        pf.first_line_indent = first_indent


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def add_inline(p, text, size=BODY_PT, color=None):
    """**粗體**、*斜體* 轉成 run。"""
    text = text.replace("<br>", " ")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_run_font(p.add_run(part[2:-2]), size, bold=True, color=color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_run_font(p.add_run(part[1:-1]), size, italic=True, color=color)
        else:
            set_run_font(p.add_run(part), size, color=color)


def shade(cell, hexcolor):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_field(p, instr, placeholder="1"):
    r = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = placeholder
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, t, s):
        r._element.append(el)
    r2 = p.add_run(); r2._element.append(txt)
    r3 = p.add_run(); r3._element.append(e)
    set_run_font(r2, SMALL_PT, color="666666")


# ---------------------------------------------------------------- 解析 Markdown → 報告區塊
def parse(md):
    lines = md.split("\n")
    blocks = []  # (kind, payload)
    i = 0
    skip_private = False
    in_figspec = None
    # 從第一個「## 封面」開始
    while i < len(lines) and not lines[i].startswith("## 封面"):
        i += 1
    while i < len(lines):
        ln = lines[i]
        if re.match(r"#{3,4} (要點|交給第 5 節)", ln):
            skip_private = True
            i += 1
            continue
        if re.match(r"#{2,4} ", ln):
            skip_private = False
            in_figspec = None
            if re.match(r"#{3,4} (正文|封面內容)$", ln):
                i += 1
                continue
            level = 1 if ln.startswith("## ") else 2
            title = re.sub(r"^#+ ", "", ln)
            title = re.sub(r"（(Appendix|Cover|References)[^）]*）$", "", title)
            blocks.append(("h", (level, title)))
            i += 1
            continue
        if skip_private or ln.startswith(">") or ln.strip() == "---":
            i += 1
            continue
        m = re.match(r"\*\*圖 (\d)｜", ln)
        if m:
            in_figspec = m.group(1)
            i += 1
            continue
        if in_figspec:
            cap = re.match(r"- \*\*圖說\*\*：(.*)", ln)
            if cap:
                blocks.append(("fig", (in_figspec, cap.group(1).strip())))
                in_figspec = None
            i += 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
            blocks.append(("table", cells))
            continue
        if ln.startswith("<p align=\"center\">"):
            blocks.append(("eq", re.sub(r"</?p[^>]*>", "", ln)))
            i += 1
            continue
        if re.match(r"\*\*表 [0-9A]+｜", ln):
            blocks.append(("tcap", ln))
            i += 1
            continue
        if (ln.startswith("註：") or ln.startswith("註 ") or re.match(r"^\**表 [0-9A]+ 註", ln)
                or (re.match(r"^[①②③④⑤]", ln) and blocks and blocks[-1][0] == "note")):
            blocks.append(("note", ln))
            i += 1
            continue
        if ln.startswith("- "):
            blocks.append(("bullet", ln[2:]))
            i += 1
            continue
        if ln.strip():
            blocks.append(("p", ln.strip()))
        i += 1
    return blocks


# ---------------------------------------------------------------- 產生文件
# 各表欄寬（cm，合計約 15.9）；長內容欄給寬、短內容欄給窄
WIDTHS = {"1": [3.4, 2.8, 4.6, 3.4, 1.7], "2": [2.6, 2.5, 2.0, 2.2, 1.9, 2.2, 2.5],
          "3": [3.0, 4.2, 6.9, 1.8], "4": [1.3, 3.4, 3.0, 4.0, 4.2], "5": [1.8, 2.4, 1.6, 2.0, 3.3, 4.8],
          "6": [1.9, 4.0, 4.2, 5.8], "7": [1.9, 2.7, 2.8, 2.6, 2.2, 3.7], "8": [1.9, 3.8, 5.6, 4.6],
          "9": [1.8, 2.8, 1.2, 1.9, 2.4, 3.4, 2.4], "10": [2.6, 1.9, 4.4, 4.6, 2.4],
          "11": [2.2, 1.2, 3.3, 5.7, 3.5], "A": [3.4, 5.6, 2.6, 4.3]}
NOTE_PT = 10
CAP_PT = 11


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def set_col_widths(t, widths):
    t.autofit = False
    for ci, w in enumerate(widths):
        if ci < len(t.columns):
            t.columns[ci].width = Cm(w)
    for row in t.rows:
        for ci, w in enumerate(widths):
            if ci < len(row.cells):
                row.cells[ci].width = Cm(w)


def build():
    from docx.enum.text import WD_TAB_ALIGNMENT

    md = SRC.read_text(encoding="utf-8")
    blocks = parse(md)
    doc = Document()
    cp = doc.core_properties
    cp.author = "黃柏凱 (Po-Kai Huang)"
    cp.last_modified_by = "Po-Kai Huang"
    cp.comments = ""
    cp.title = "把電話與優惠留給打了才會買的人"
    cp.subject = "321513 Machine Learning - Assessment Task 2"

    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))

    style_fonts(doc.styles["Normal"], BODY_PT)
    op = OxmlElement("w:overflowPunct")  # 行尾標點不突出版心
    op.set(qn("w:val"), "0")
    doc.styles["Normal"].element.get_or_add_pPr().append(op)
    style_fonts(doc.styles["Heading 1"], 14, bold=True)
    style_fonts(doc.styles["Heading 2"], 13, bold=True)
    for name in ("Heading 1", "Heading 2"):
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(12), Pt(6), 1.5
        pf.keep_with_next = True

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE", "1")

    # ---- 封面
    assert blocks[0] == ("h", (1, "封面"))
    idx = 1
    cover_lines, cover_table = [], None
    while blocks[idx][0] != "h":
        k, v = blocks[idx]
        if k == "p":
            cover_lines.append(v)
        elif k == "table":
            cover_table = v
        idx += 1
    for _ in range(7):
        para_format(doc.add_paragraph(), 1.0)
    for j, t in enumerate(cover_lines):
        p = doc.add_paragraph()
        para_format(p, 1.3, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_inline(p, t, size=(20 if j == 0 else 15 if j == 1 else 12))
    para_format(doc.add_paragraph(), 1.0, after=30)
    tbl = doc.add_table(rows=0, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in cover_table[1:]:
        r = tbl.add_row().cells
        for c, txt in zip(r, row):
            c.text = ""
            p = c.paragraphs[0]
            para_format(p, 1.3, after=4)
            add_inline(p, txt, size=11)
        set_run_font(r[0].paragraphs[0].runs[0], 11, bold=True)
    set_col_widths(tbl, [3.2, 12.0])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- 目錄
    p = doc.add_paragraph()
    para_format(p, 1.5, after=12)
    add_inline(p, "**目錄**", size=14)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "（目錄）")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- 內文
    in_refs = False
    cur_table = None
    pending_note = None  # 表標題括號內的說明，移到表格下方
    for k, v in blocks[idx:]:
        if k == "h":
            level, title = v
            if title.startswith("參考文獻") or title.startswith("附錄 A"):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            in_refs = title.startswith("參考文獻")
            h = doc.add_heading(level=level)
            add_inline(h, title, size=14 if level == 1 else 13)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
        elif k == "p":
            p = doc.add_paragraph()
            if in_refs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.line_spacing = Pt(30)
                p.paragraph_format.left_indent = Cm(1.27)
                p.paragraph_format.first_line_indent = Cm(-1.27)
                add_inline(p, v, size=BODY_PT)
            else:
                para_format(p, 2.0, after=0)
                add_inline(p, v)
        elif k == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            para_format(p, 2.0, after=0)
            add_inline(p, v, size=BODY_PT)
        elif k == "eq":
            p = doc.add_paragraph()
            parts = [x for x in re.split("　{2,}", v) if x]
            if len(parts) == 2:
                para_format(p, 1.5, before=4, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
                ts = p.paragraph_format.tab_stops
                ts.add_tab_stop(Cm(7.96), WD_TAB_ALIGNMENT.CENTER)
                ts.add_tab_stop(Cm(15.92), WD_TAB_ALIGNMENT.RIGHT)
                add_inline(p, "\t" + parts[0] + "\t" + parts[1])
            else:
                para_format(p, 1.5, before=4, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
                add_inline(p, v)
        elif k == "tcap":
            m = re.match(r"(\*\*表 ([0-9A]+)｜[^*]+\*\*)(.*)$", v)
            cur_table = m.group(2) if m else None
            title, rest = (m.group(1), m.group(3).strip()) if m else (v, "")
            if rest.startswith("（") and rest.endswith("）") and cur_table != "3":
                pending_note = "註：" + rest[1:-1]
                rest = ""
            p = doc.add_paragraph()
            para_format(p, 1.2, before=10, after=4)
            p.paragraph_format.keep_with_next = True
            add_inline(p, title + (" " + rest if rest else ""), size=CAP_PT)
        elif k == "note":
            p = doc.add_paragraph()
            para_format(p, 1.2, before=3, after=6)
            p.paragraph_format.keep_together = True
            add_inline(p, v, size=NOTE_PT, color="333333")
        elif k == "table":
            ncol = len(v[0])
            t = doc.add_table(rows=0, cols=ncol)
            t.style = doc.styles["Table Grid"]
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(v):
                cells = t.add_row().cells
                no_split(t.rows[-1])
                row = (row + [""] * ncol)[:ncol]
                for c, txt in zip(cells, row):
                    c.text = ""
                    split = ri > 0 and (txt.count("**") >= 4 or re.search("｜[^｜]{1,6}：", txt))
                    parts = [s.strip() for s in txt.split("｜")] if split else [txt]
                    for pi, part in enumerate(parts):
                        p = c.paragraphs[0] if pi == 0 else c.add_paragraph()
                        para_format(p, 1.1, after=1)
                        if ri == 0:
                            p.paragraph_format.keep_with_next = True
                        add_inline(p, part, size=9 if ncol < 6 else 8.5)
                    if ri == 0:
                        shade(c, "E8E6E1")
                        for r in c.paragraphs[0].runs:
                            r.bold = True
                if ri == 0:
                    repeat_header(t.rows[0])
            if cur_table in WIDTHS and len(WIDTHS[cur_table]) == ncol:
                set_col_widths(t, WIDTHS[cur_table])
            cur_table = None
            if pending_note:
                p = doc.add_paragraph()
                para_format(p, 1.2, before=3, after=6)
                add_inline(p, pending_note, size=NOTE_PT, color="333333")
                pending_note = None
            else:
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif k == "fig":
            num, cap = v
            p = doc.add_paragraph()
            para_format(p, 1.0, before=8, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(FIGS[num]), width=Cm(15.8))
            c = doc.add_paragraph()
            para_format(c, 1.2, after=10)
            add_inline(c, cap.replace(f"圖 {num}｜", f"**圖 {num}｜**", 1), size=CAP_PT)
    doc.save(OUT)
    return OUT, blocks


if __name__ == "__main__":
    out, blocks = build()
    kinds = {}
    for k, _ in blocks:
        kinds[k] = kinds.get(k, 0) + 1
    print(out, kinds)
